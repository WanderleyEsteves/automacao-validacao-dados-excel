import pandas as pd
import os, sys, re, threading
from openpyxl import load_workbook
from openpyxl.styles import PatternFill
import tkinter as tk
from tkinter import messagebox
import gspread
from oauth2client.service_account import ServiceAccountCredentials

# Configurações de arquivo e API do Google
NOME_ARQUIVO_JSON = ""  # Altere para o nome do seu arquivo .json
ID_PLANILHA_GOOGLE = ""      # Coloque aqui o ID da sua planilha (o código entre /d/ e /edit)
arquivo_excel = "definitivo_2.xlsx"

def caminho_recurso(relative_path):
    try:
        
        base_path = sys._MEIPASS
    except Exception:
        
        try:
            base_path = os.path.dirname(os.path.abspath(__file__))
        except NameError:
            base_path = os.path.abspath(".")
            
    return os.path.join(base_path, relative_path)


def formatar_cpf(val):
    numeros = re.sub(r"\D", "", str(val))
    if len(numeros) != 11:
        return val
    return f"{numeros[:3]}.{numeros[3:6]}.{numeros[6:9]}-{numeros[9:]}"

def cpf_valido(cpf):
    cpf = re.sub(r"\D", "", str(cpf))
    if len(cpf) != 11 or cpf == cpf[0] * 11: return False
    for peso in (10, 11):
        soma = sum(int(cpf[i]) * (peso - i) for i in range(peso - 1))
        resto = (soma * 10) % 11
        if (resto if resto < 10 else 0) != int(cpf[peso - 1]): return False
    return True

def tratar_e_formatar_telefone(val):
    numeros = re.sub(r"\D", "", str(val))
    if len(numeros) == 13 and numeros.startswith("55"):
        numeros = numeros[2:]
    if len(numeros) == 11:
        return f"({numeros[:2]}) {numeros[2:7]}-{numeros[7:]}", True
    return val, False

def conectar_google_sheets():

    scopes = [
        "https://spreadsheets.google.com/feeds",
        "https://www.googleapis.com/auth/drive"
    ]
    caminho_json = caminho_recurso(NOME_ARQUIVO_JSON)
    creds = ServiceAccountCredentials.from_json_keyfile_name(caminho_json, scopes)
    client = gspread.authorize(creds)
    sheet = client.open_by_key(ID_PLANILHA_GOOGLE).sheet1
    return sheet

def executar_atualizacao():
    try:
        label_status.config(text="Buscando novas respostas...", fg="#FFFFFF")
        btn_atualizar.config(state="disabled")
        
        sheet = conectar_google_sheets()
        dados_brutos = sheet.get_all_values()
        
        if dados_brutos:
            headers = dados_brutos[0]
            rows = dados_brutos[1:]
            df = pd.DataFrame(rows, columns=headers).astype(str).fillna("")
        else:
            df = pd.DataFrame()

        if "CPF" in df:
            df["CPF"] = df["CPF"].apply(formatar_cpf)

        termos_telefone = ["TELEFONE", "CELULAR","TELEFONE (ZAP)" ,"CEL", "FONE", "WHATS", "CONTATO"]
        coluna_tel = next((col for col in df.columns if any(t in col.strip().upper() for t in termos_telefone)), None)

        status_cpfs = [not cpf_valido(x) for x in df["CPF"]] if "CPF" in df else []
        status_tels = []

        if coluna_tel:
            for index, row in df.iterrows():
                tel_formatado, valido = tratar_e_formatar_telefone(row[coluna_tel])
                df.at[index, coluna_tel] = tel_formatado
                status_tels.append(not valido)

        label_status.config(text="Salvando na planilha...", fg="#FFFFFF")
        df.to_excel(arquivo_excel, index=False)

        wb = load_workbook(arquivo_excel)
        ws = wb.active
        
        cols = [ws.cell(row=1, column=c).value for c in range(1, ws.max_column + 1)]
        idx_cpf = cols.index("CPF") + 1 if "CPF" in cols else None
        
        idx_tel = next((cols.index(c)+1 for c in cols if c and any(t in c.strip().upper() for t in termos_telefone)), None)

        vermelho = PatternFill(start_color="FFC7CE", end_color="FFC7CE", fill_type="solid")
        laranja = PatternFill(start_color="FFDAB9", end_color="FFDAB9", fill_type="solid")

        for idx, (cpf_errado, tel_errado) in enumerate(zip(status_cpfs, status_tels), start=2):
            if idx_cpf: ws.cell(row=idx, column=idx_cpf).fill = vermelho if cpf_errado else PatternFill(fill_type=None)
            if idx_tel: ws.cell(row=idx, column=idx_tel).fill = laranja if tel_errado else PatternFill(fill_type=None)

        wb.save(arquivo_excel)
        label_status.config(text="Pronto! Tudo atualizado.", fg="#4ADE80")
        messagebox.showinfo("Sucesso", "A planilha foi atualizada com os novos dados.")

    except FileNotFoundError:
        label_status.config(text="Erro de Arquivo.", fg="#F87171")
        messagebox.showerror("Arquivo não encontrado", f"Não foi possível localizar o arquivo de chave '{NOME_ARQUIVO_JSON}'.")
    except PermissionError:
        label_status.config(text="Erro: Arquivo aberto.", fg="#F87171")
        messagebox.showerror("Arquivo em uso", f"A planilha '{arquivo_excel}' está aberta. Por favor, feche-a para continuar.")
    except Exception as e:
        label_status.config(text="Ops, algo deu errado.", fg="#F87171")
        messagebox.showerror("Erro", f"Não foi possível atualizar: {str(e)}")
    finally:
        btn_atualizar.config(state="normal")

# INTERFACE 

janela = tk.Tk()
janela.title("x - CONTRATO x x")
janela.geometry("420x450")
janela.configure(bg="#FFFFFF")
janela.resizable(False, False)

# Logo
try:
    caminho_logo = caminho_recurso("logo.png")
    img = tk.PhotoImage(file=caminho_logo)
    if img.width() > 250: img = img.subsample(2, 2)
    tk.Label(janela, image=img, bg="#FCFDFF").pack(pady=15)
except Exception:
    tk.Label(janela, text="SISTEMA X", font=("Arial", 20, "bold"), bg="#FFFFFF", fg="#FFFFFF").pack(pady=40)

tk.Label(janela, text="Clique abaixo para buscar \n dados usuários.", font=("Arial", 11), bg="#FFFFFF", fg="#1A2B4C").pack(pady=10)

btn_atualizar = tk.Button(janela, text="ATUALIZAR PLANILHA", font=("Arial", 11, "bold"), bg="#D1D5DB", fg="#1A2B4C", relief="flat", cursor="hand2", padx=60,pady=10, command=lambda: threading.Thread(target=executar_atualizacao, daemon=True).start())
btn_atualizar.pack(pady=20)

label_status = tk.Label(janela, text="Pronto para usar.", font=("Arial", 10, "italic"), bg="#FFFFFF", fg="#FFFFFF")
label_status.pack(pady=10)

janela.mainloop()